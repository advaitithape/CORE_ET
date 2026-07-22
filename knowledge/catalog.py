"""Document catalog + incremental ingestion.

The knowledge base starts EMPTY and grows as documents are ingested.
Two ingestion paths:
  - STRUCTURED: recognised industrial documents (PFD/PFMEA/Control Plan/QMS records)
    contribute a slice of their part's canonical data to the graph.
  - GENERAL: any other readable document (reports, notes, essays, finance sheets…)
    has its text extracted, chunked and embedded so the assistant can answer from it
    with citations — nothing uploadable breaks the system.
"""
import os
from .base import load_parts, build_graph_from_parts, find_file

# ---- known structured demo documents: basename -> (part_no, doc_type) -------
KNOWN_DOCS = {
    # Part 1 — CE21609 (real files)
    "CE21609_ Process Flow Chart 2.xls": ("CE21609", "pfd"),
    "CE21609_ PFMEA 2.xls": ("CE21609", "pfmea"),
    "CE21609_CONTROL PLAN 2.xlsx": ("CE21609", "control_plan"),
    "CE21609 1.pdf": ("CE21609", "drawing"),
    # Part 2 — S-10254-5 Base
    "S-10254-5_PFD.xlsx": ("S-10254-5", "pfd"),
    "S-10254-5_PFMEA.xlsx": ("S-10254-5", "pfmea"),
    "S-10254-5_ControlPlan.xlsx": ("S-10254-5", "control_plan"),
    "S-10254-5_InwardInspection.xlsx": ("S-10254-5", "inward_inspection"),
    "S-10254-5_InProcessIR.xlsx": ("S-10254-5", "in_process_ir"),
    "BASE IR.xlsx": ("S-10254-5", "in_process_ir"),
    "S-10254-5_PDIR.xlsx": ("S-10254-5", "pdir"),
    "S-10254-5_NCR_CAR.xlsx": ("S-10254-5", "ncr_car"),
    "S-10254-5_8D_NCR-BASE-014.xlsx": ("S-10254-5", "capa_8d"),
    "Work Instruction SI-L-06 Hindi1.docx": ("S-10254-5", "work_instruction"),
    # Part 3 — YF-4021 Yoke Flange
    "YF-4021_PFD.xlsx": ("YF-4021", "pfd"),
    "YF-4021_PFMEA.xlsx": ("YF-4021", "pfmea"),
    "YF-4021_ControlPlan.xlsx": ("YF-4021", "control_plan"),
    "YF-4021_InwardInspection.xlsx": ("YF-4021", "inward_inspection"),
    "YF-4021_InProcessIR.xlsx": ("YF-4021", "in_process_ir"),
    "YF-4021_PDIR.xlsx": ("YF-4021", "pdir"),
    "YF-4021_NCR_CAR.xlsx": ("YF-4021", "ncr_car"),
    "YF-4021_8D_NCR-YF-007.xlsx": ("YF-4021", "capa_8d"),
    "YF-4021_WorkInstruction.docx": ("YF-4021", "work_instruction"),
}


# ---- generic text extraction (any readable document) ------------------------

def extract_text(path):
    """Best-effort plain-text extraction for arbitrary uploads."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".txt", ".md", ".csv", ".log", ".json"):
            with open(path, encoding="utf-8", errors="replace") as f:
                return f.read()
        if ext == ".docx":
            from docx import Document
            return "\n".join(p.text for p in Document(path).paragraphs)
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                return "\n".join((pg.extract_text() or "") for pg in pdf.pages[:40])
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            out = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v is not None and str(v).strip()]
                    if vals:
                        out.append(" | ".join(vals))
            return "\n".join(out)
        if ext == ".xls":
            import xlrd
            wb = xlrd.open_workbook(path)
            out = []
            for sh in wb.sheets():
                for r in range(sh.nrows):
                    vals = [str(sh.cell_value(r, c)) for c in range(sh.ncols)
                            if str(sh.cell_value(r, c)).strip()]
                    if vals:
                        out.append(" | ".join(vals))
            return "\n".join(out)
    except Exception:  # noqa: BLE001 — unreadable upload must not break ingestion
        return ""
    return ""


def chunk_text(text, size=900, cap=60):
    paras = [p.strip() for p in text.split("\n") if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        if len(chunks) >= cap:
            break
        if cur and len(cur) + len(p) + 1 > size:
            chunks.append(cur)
            cur = p
        else:
            cur = (cur + "\n" + p).strip()
    if cur and len(chunks) < cap:
        chunks.append(cur)
    return chunks


def _op_shell(o):
    return {"op_no": o["op_no"], "name": o.get("name", ""), "machine": o.get("machine", ""),
            "cp_ref": o.get("cp_ref"), "product_chars": [], "process_chars": []}


def build_from_ingested(basenames):
    """Assemble the KB from only the ingested documents (basenames)."""
    fulls = {p["meta"]["part_no"]: p for p in load_parts()}
    acc = {}
    general = []   # [(basename, n_chunks)]

    def part(pno):
        return acc.setdefault(pno, {"meta": fulls[pno]["meta"], "operations": {},
                                    "failure_modes": [], "qms": {}, "docs": set()})

    general_chunks = []
    for bn in basenames:
        info = KNOWN_DOCS.get(bn)
        if not info:
            # ---- GENERAL document: extract + chunk its text ----
            path = find_file(bn)
            text = extract_text(path) if path else ""
            chunks = chunk_text(text)
            for i, c in enumerate(chunks):
                general_chunks.append({"id": f"gen:{bn}:{i}",
                                       "text": f"[{bn}] {c}",
                                       "meta": {"part_no": None, "doc_type": "general",
                                                "file": bn, "op_no": None}})
            general.append((bn, len(chunks)))
            continue
        pno, dtype = info
        if pno not in fulls:
            continue   # structured source files not present (e.g. CE21609 not yet uploaded)
        full = fulls[pno]
        a = part(pno)
        a["docs"].add(dtype)
        if dtype == "pfd":
            for o in full["operations"]:
                a["operations"].setdefault(o["op_no"], _op_shell(o))
        elif dtype == "control_plan":
            for o in full["operations"]:
                op = a["operations"].setdefault(o["op_no"], _op_shell(o))
                op["product_chars"] = o.get("product_chars", [])
                op["process_chars"] = o.get("process_chars", [])
        elif dtype == "pfmea":
            a["failure_modes"] = full["failure_modes"]
            for f in full["failure_modes"]:
                if f.get("op_no"):
                    a["operations"].setdefault(f["op_no"], {"op_no": f["op_no"], "name": f.get("function", ""),
                        "machine": "", "cp_ref": None, "product_chars": [], "process_chars": []})
        elif dtype == "inward_inspection":
            a["qms"]["inward"] = full["qms"].get("inward")
        elif dtype == "in_process_ir":
            a["qms"]["inprocess"] = full["qms"].get("inprocess")
        elif dtype == "pdir":
            a["qms"]["pdir"] = full["qms"].get("pdir")
        elif dtype == "ncr_car":
            a["qms"]["ncrs"] = full["qms"].get("ncrs")
        elif dtype == "capa_8d":
            a["qms"]["capas"] = full["qms"].get("capas")
        elif dtype == "work_instruction":
            a["qms"]["wi"] = full["qms"].get("wi")
        # 'drawing' currently contributes no structural data (CV roadmap)

    parts = [{"meta": a["meta"], "operations": list(a["operations"].values()),
              "failure_modes": a["failure_modes"], "qms": a["qms"]}
             for a in acc.values()]
    kb = build_graph_from_parts(parts)
    kb["chunks"].extend(general_chunks)
    kb["general_docs"] = [{"name": bn, "chunks": n} for bn, n in general]
    kb["stats"]["general_docs"] = len(general)
    kb["stats"]["chunks"] = len(kb["chunks"])
    kb["ingested"] = list(basenames)
    return kb


def is_ingestable(basename):
    """Structured demo document, or any upload we can extract text from."""
    if basename in KNOWN_DOCS:
        return True
    p = find_file(basename)
    return bool(p and extract_text(p).strip())


def ingest_kind(basename):
    if basename in KNOWN_DOCS:
        return "structured"
    p = find_file(basename)
    if p and extract_text(p).strip():
        return "general"
    return None
